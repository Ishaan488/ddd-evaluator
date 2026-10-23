import openpyxl
import pytest

from ddd_evaluator.scorecard_io import (
    DIM_SHEET, ScorecardFileError, content_hash, read_scorecard, template_card, write_scorecard_xlsx, write_template,
)
from conftest import SCORECARD, edit_scorecard


def problems_of(path):
    with pytest.raises(ScorecardFileError) as exc:
        read_scorecard(path)
    return exc.value.problems


def set_cell(sheet, row, column_name, value):
    def mutate(wb):
        ws = wb[sheet]
        header = [c.value for c in ws[1]]
        ws.cell(row=row, column=header.index(column_name) + 1).value = value
    return mutate


def test_the_real_scorecard_reads_cleanly():
    card = read_scorecard(SCORECARD)
    assert len(card.dimensions) == 19 and card.frozen
    assert card.thresholds.green == 90 and card.thresholds.critical_min == 4
    assert card.get("L02").anchors[5].startswith("Every phrase")
    assert card.get("L03").look_for  # one item per line survived the trip through Excel


def test_write_then_read_gives_the_same_scorecard(tmp_path):
    card = read_scorecard(SCORECARD)
    out = tmp_path / "copy.xlsx"
    write_scorecard_xlsx(card, out)
    assert read_scorecard(out) == card


def test_content_hash_ignores_resaving_but_notices_a_change(tmp_path):
    a = read_scorecard(SCORECARD)
    resaved = edit_scorecard(tmp_path, lambda wb: None, "resaved.xlsx")
    assert content_hash(read_scorecard(resaved)) == content_hash(a)
    changed = edit_scorecard(tmp_path, set_cell(DIM_SHEET, 3, "Name", "A different name"), "changed.xlsx")
    assert content_hash(read_scorecard(changed)) != content_hash(a)


def test_template_is_a_valid_but_unfrozen_example(tmp_path):
    out = tmp_path / "t.xlsx"
    write_template(out)
    card = read_scorecard(out)
    assert card.frozen is False and sum(d.weight for d in card.dimensions) == 100
    assert {d.evaluator for d in card.dimensions} == {"code", "llm", "human"}
    assert template_card().id == "SC-EXAMPLE"
    wb = openpyxl.load_workbook(out)
    assert {"Scorecard", "Dimensions", "Instructions", "Code checks"} <= set(wb.sheetnames)


def test_all_problems_are_reported_at_once_with_their_rows(tmp_path):
    def mutate(wb):
        ws = wb[DIM_SHEET]
        header = [c.value for c in ws[1]]
        col = lambda n: header.index(n) + 1
        ws.cell(row=2, column=col("Weight")).value = "heavy"          # not a number
        ws.cell(row=3, column=col("Scored by")).value = "robot"       # not code/llm/human
        ws.cell(row=4, column=col("Anchor 3")).value = None           # missing anchor
        ws.cell(row=5, column=col("Checks")).value = "K99"            # unknown check
        ws.cell(row=10, column=col("Guidance")).value = None          # llm without guidance
    p = problems_of(edit_scorecard(tmp_path, mutate))
    text = "\n".join(p)
    assert "row 2" in text and "Weight must be a number" in text
    assert "row 3" in text and "'Scored by' must be code, llm or human" in text
    assert "row 4" in text and "Anchor 3 is empty" in text
    assert "row 5" in text and "unknown check(s) K99" in text
    assert "row 10" in text and "needs Guidance" in text
    assert len(p) >= 5


def test_duplicate_ids_and_wrong_total(tmp_path):
    p = problems_of(edit_scorecard(tmp_path, set_cell(DIM_SHEET, 3, "ID", "D01")))
    assert any("'D01' is used more than once" in x for x in p)
    p = problems_of(edit_scorecard(tmp_path, set_cell(DIM_SHEET, 2, "Weight", 9), "w.xlsx"))
    assert any("weights add up to 104" in x for x in p)


def test_missing_sheets_and_columns_and_settings(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "Other"
    bare = tmp_path / "bare.xlsx"
    wb.save(bare)
    assert "Sheet(s) missing" in problems_of(bare)[0]

    def drop_anchor_column(wb):
        ws = wb[DIM_SHEET]
        header = [c.value for c in ws[1]]
        ws.delete_cols(header.index("Anchor 2") + 1)
    assert "column(s) missing: Anchor 2" in problems_of(edit_scorecard(tmp_path, drop_anchor_column, "c.xlsx"))[0]

    def clear_id(wb):
        for row in wb["Scorecard"].iter_rows(min_row=1, max_col=2):
            if row[0].value == "Scorecard ID":
                row[1].value = None
    assert any("'Scorecard ID' is empty" in x for x in problems_of(edit_scorecard(tmp_path, clear_id, "d.xlsx")))


def test_checks_are_only_for_code_dimensions(tmp_path):
    p = problems_of(edit_scorecard(tmp_path, set_cell(DIM_SHEET, 7, "Checks", "K01"), "e.xlsx"))
    assert any("only for code dimensions" in x for x in p)


def test_not_an_excel_file(tmp_path):
    junk = tmp_path / "junk.xlsx"
    junk.write_text("not a workbook")
    assert "could not be opened" in problems_of(junk)[0]
