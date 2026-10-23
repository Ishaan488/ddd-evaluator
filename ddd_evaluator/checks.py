"""Deterministic checks on the artifact workbook (Tree, Schema register, Allowed values).

These run in plain code, never through an LLM. Each check returns how many items were
examined and which of them broke the rule, so the dimension score is reproducible.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Optional

from .models import Dimension, Severity

DATA_TYPES = {"id", "text", "enum", "integer", "decimal", "date", "reference", "boolean", "timestamp"}
REQUIRED_VALUES = {"required", "optional", "conditional"}

# Reserved in PostgreSQL or MySQL 8 (the docket asks for both engines to be checked).
RESERVED_WORDS = {
    "analyse", "analyze", "array", "asc", "both", "case", "check", "column", "condition", "default",
    "desc", "end", "from", "grant", "group", "groups", "index", "key", "lateral", "leading", "limit",
    "natural", "offset", "only", "order", "outer", "over", "partition", "precision", "primary", "range",
    "rank", "read", "references", "returning", "row", "rows", "schema", "select", "some", "symmetric",
    "system", "table", "trailing", "trigger", "unique", "user", "using", "variadic", "verbose", "when",
    "where", "window", "with", "write",
}
SINGULAR_EXCEPTIONS = {"status", "address", "class", "series", "analysis", "access", "process"}

DATA_SHEETS = {  # register table -> (sheet, key column that marks the header row)
    "nfr": ("NFR register", "nfr_id"),
    "requirement": ("Requirements", "requirement_id"),
    "scope_line": ("Scope lines", "scope_id"),
}


@dataclass
class CheckResult:
    id: str
    name: str
    severity: Severity
    total: int
    violations: list[str]
    fix: str
    note: str = ""

    @property
    def passed(self) -> bool:
        return not self.violations

    @property
    def pass_rate(self) -> float:
        if self.total <= 0:
            return 1.0 if not self.violations else 0.0
        return max(0.0, 1.0 - len(self.violations) / self.total)

    def summary(self) -> str:
        status = "PASS" if self.passed else f"FAIL ({len(self.violations)} of {self.total})"
        return f"{self.id} {self.name}: {status}"


# --------------------------------------------------------------------------- workbook access
def _rows(ws, key: Optional[str] = None) -> list[dict]:
    raw = list(ws.iter_rows(values_only=True))
    if not raw:
        return []
    start = 0
    if key is not None:
        for i, row in enumerate(raw):
            if key in row:
                start = i
                break
        else:
            return []
    header = [str(h).strip() if h is not None else "" for h in raw[start]]
    out = []
    for r in raw[start + 1 :]:
        if not any(c is not None and str(c).strip() for c in r):
            continue
        out.append({h: ("" if v is None else str(v).strip()) for h, v in zip(header, r) if h})
    return out


@dataclass
class ArtifactWorkbook:
    tree: list[dict]
    register: list[dict]
    allowed: list[dict]
    data: dict[str, list[dict]] = field(default_factory=dict)

    @classmethod
    def load(cls, path: Path) -> Optional["ArtifactWorkbook"]:
        import openpyxl

        if path.suffix.lower() not in {".xlsx", ".xlsm"}:
            return None
        wb = openpyxl.load_workbook(path, data_only=True)
        need = ["Tree", "Schema register", "Allowed values"]
        if not all(n in wb.sheetnames for n in need):
            return None
        data = {}
        for table, (sheet, key) in DATA_SHEETS.items():
            if sheet in wb.sheetnames:
                data[table] = _rows(wb[sheet], key)
        return cls(_rows(wb["Tree"]), _rows(wb["Schema register"]), _rows(wb["Allowed values"]), data)

    # helpers
    def level3(self) -> list[dict]:
        return [n for n in self.tree if _level(n) == 3]

    def list_values(self) -> dict[str, set[str]]:
        out: dict[str, set[str]] = {}
        for r in self.allowed:
            if r.get("List"):
                out.setdefault(r["List"], set()).add(r.get("Value", ""))
        return out


def _level(node: dict) -> int:
    try:
        return int(float(node.get("Level", "0")))
    except ValueError:
        return 0


def _list_ref(rule: str) -> Optional[str]:
    m = re.search(r"list:\s*([A-Za-z0-9_]+)", rule or "")
    return m.group(1) if m else None


def _real_column(col: str) -> bool:
    return bool(col) and not col.startswith("(")


def _head(items: list[str], n: int = 6) -> list[str]:
    return items[:n] + ([f"... and {len(items) - n} more"] if len(items) > n else [])


# --------------------------------------------------------------------------- the checks
def k01(wb: ArtifactWorkbook) -> CheckResult:
    bad = [f"{n.get('Path id')} (level {_level(n)})" for n in wb.tree if _level(n) > 3]
    return CheckResult("K01", "Tree is no deeper than three levels", "HIGH", len(wb.tree), bad,
                       "Flatten any node deeper than level 3 into a sub-parameter.")


def k02(wb: ArtifactWorkbook) -> CheckResult:
    ids = [n.get("Path id", "") for n in wb.tree]
    dupes = sorted({i for i in ids if ids.count(i) > 1})
    return CheckResult("K02", "Every node has its own path id", "HIGH", len(ids), dupes,
                       "Give each repeated path id a unique value.")


def k03(wb: ArtifactWorkbook) -> CheckResult:
    bad = [f"{n.get('Path id')} {n.get('Name')}" for n in wb.tree if not n.get("Definition")]
    return CheckResult("K03", "Every node has a definition", "HIGH", len(wb.tree), bad,
                       "Write a one-sentence definition for each listed node.")


def k04(wb: ArtifactWorkbook) -> CheckResult:
    l3 = wb.level3()
    bad = [f"{n.get('Path id')} {n.get('Name')}" for n in l3 if not n.get("Data type")]
    return CheckResult("K04", "Every sub-parameter has a data type", "HIGH", len(l3), bad,
                       "Give each listed sub-parameter a data type.")


def k05(wb: ArtifactWorkbook) -> CheckResult:
    l3 = wb.level3()
    bad = [f"{n.get('Path id')} {n.get('Name')}" for n in l3 if not _real_column(n.get("Column", ""))]
    return CheckResult("K05", "Every sub-parameter maps to a table column", "HIGH", len(l3), bad,
                       "Name the table column for each listed sub-parameter, or remove the sub-parameter.")


def k06(wb: ArtifactWorkbook) -> CheckResult:
    lists = wb.list_values()
    enums = [n for n in wb.level3() if n.get("Data type") == "enum"]
    bad = []
    for n in enums:
        ref = _list_ref(n.get("Allowed values", ""))
        if ref is None:
            bad.append(f"{n.get('Path id')} {n.get('Name')}: no 'list:' reference")
        elif ref not in lists:
            bad.append(f"{n.get('Path id')} {n.get('Name')}: list '{ref}' is not in Allowed values")
    return CheckResult("K06", "Every choice field points to an allowed-value list that exists", "HIGH", len(enums), bad,
                       "Add the missing list to Allowed values or correct the list name.")


def k07(wb: ArtifactWorkbook) -> CheckResult:
    bad = [f"{r.get('Table')}.{r.get('Column')}" for r in wb.register if not r.get("Used by id") and not r.get("Key")]
    return CheckResult("K07", "No column exists without a sub-parameter or a key", "MEDIUM", len(wb.register), bad,
                       "Tie each listed column to a sub-parameter or mark it as a key.")


def k08(wb: ArtifactWorkbook) -> CheckResult:
    index = {n.get("Path id"): n for n in wb.tree}
    bad = []
    total = 0
    for r in wb.register:
        pid = r.get("Used by id")
        if not pid:
            continue
        total += 1
        node = index.get(pid)
        if node is None:
            bad.append(f"{r.get('Table')}.{r.get('Column')}: path id {pid} is not in the Tree")
        elif node.get("Table") != r.get("Table") or node.get("Column") != r.get("Column"):
            bad.append(f"{r.get('Table')}.{r.get('Column')}: Tree {pid} says {node.get('Table')}.{node.get('Column')}")
    return CheckResult("K08", "Each register row matches the Tree node it names", "HIGH", total, bad,
                       "Make the table and column in the register and the Tree agree for the listed rows.")


def k09(wb: ArtifactWorkbook) -> CheckResult:
    in_register = {(r.get("Table"), r.get("Column")) for r in wb.register}
    l3 = [n for n in wb.level3() if n.get("Table") and _real_column(n.get("Column", ""))]
    bad = [f"{n.get('Path id')} {n.get('Table')}.{n.get('Column')}" for n in l3
           if (n.get("Table"), n.get("Column")) not in in_register]
    return CheckResult("K09", "Every Tree column exists in the register", "HIGH", len(l3), bad,
                       "Add the missing columns to the Schema register.")


def k10(wb: ArtifactWorkbook) -> CheckResult:
    columns = {(r.get("Table"), r.get("Column")) for r in wb.register}
    fks = [r for r in wb.register if "foreign key" in r.get("Key", "").lower()]
    bad = []
    for r in fks:
        m = re.search(r"foreign key to\s+(\w+)\.(\w+)", r.get("Key", ""), re.I)
        name = f"{r.get('Table')}.{r.get('Column')}"
        if not m:
            bad.append(f"{name}: foreign key with no stated target table.column")
        elif (m.group(1), m.group(2)) not in columns:
            bad.append(f"{name}: target {m.group(1)}.{m.group(2)} is not in the register")
    return CheckResult("K10", "Every foreign key names a target that exists", "HIGH", len(fks), bad,
                       "State 'Foreign key to table.column' for each listed key.")


def k11(wb: ArtifactWorkbook) -> CheckResult:
    tables = sorted({r.get("Table") for r in wb.register if r.get("Table")})
    with_pk = {r.get("Table") for r in wb.register if "primary key" in r.get("Key", "").lower()}
    bad = [t for t in tables if t not in with_pk]
    return CheckResult("K11", "Every table has a primary key", "HIGH", len(tables), bad,
                       "Define a primary key for each listed table.")


def k12(wb: ArtifactWorkbook) -> CheckResult:
    cond = [r for r in wb.register if r.get("Required", "").lower() == "conditional"]
    stated = re.compile(r"required (when|unless|if)\b|only (when|if)\b", re.I)
    bad = [f"{r.get('Table')}.{r.get('Column')}: rule is '{r.get('Allowed values')}'" for r in cond
           if not stated.search(r.get("Allowed values", ""))]
    return CheckResult("K12", "Every conditional column states its condition", "MEDIUM", len(cond), bad,
                       "Write 'required when ...' in the rule of each listed column.")


def k13(wb: ArtifactWorkbook) -> CheckResult:
    bad = []
    for r in wb.register:
        name = f"{r.get('Table')}.{r.get('Column')}"
        if r.get("Required", "").lower() not in REQUIRED_VALUES:
            bad.append(f"{name}: Required = '{r.get('Required')}'")
        if r.get("Data type", "").lower() not in DATA_TYPES:
            bad.append(f"{name}: Data type = '{r.get('Data type')}'")
    return CheckResult("K13", "Required flags and data types use the agreed vocabulary", "MEDIUM", len(wb.register) * 2, bad,
                       "Use only required/optional/conditional and the agreed data types.")


def k14(wb: ArtifactWorkbook) -> CheckResult:
    pks = [r for r in wb.register if r.get("Data type") == "id" and r.get("Key", "").lower().startswith("primary key")]
    bad = [f"{r.get('Table')}.{r.get('Column')}" for r in pks if "pattern" not in r.get("Allowed values", "").lower()]
    return CheckResult("K14", "Every id key has a stated pattern", "LOW", len(pks), bad,
                       "State the id pattern (for example FD-nnn) for each listed key.")


def k15(wb: ArtifactWorkbook) -> CheckResult:
    index = {n.get("Path id"): n for n in wb.tree}
    bad = []
    total = 0
    for r in wb.register:
        pid = r.get("Used by id")
        node = index.get(pid) if pid else None
        if node is None:
            continue
        total += 1
        register_derived = r.get("L4 atomic definition", "").lower().startswith("derived:")
        tree_derived = "deriv" in node.get("Atomic or derived", "").lower()
        if register_derived != tree_derived:
            bad.append(f"{r.get('Table')}.{r.get('Column')}: register says "
                       f"{'derived' if register_derived else 'atomic'}, Tree says '{node.get('Atomic or derived')}'")
    return CheckResult("K15", "Derived columns are labelled the same way in Tree and register", "MEDIUM", total, bad,
                       "Make the atomic/derived label agree in both sheets.")


def k16(wb: ArtifactWorkbook) -> CheckResult:
    lists = wb.list_values()
    used = set()
    for n in wb.level3():
        ref = _list_ref(n.get("Allowed values", ""))
        if ref:
            used.add(ref)
    for r in wb.register:
        ref = _list_ref(r.get("Allowed values", ""))
        if ref:
            used.add(ref)
    bad = [f"list {name} is not used by any column" for name in sorted(lists) if name not in used]
    seen = set()
    for r in wb.allowed:
        key = (r.get("List"), r.get("Value"))
        if not r.get("Meaning"):
            bad.append(f"{key[0]} / {key[1]}: no meaning")
        if key in seen:
            bad.append(f"{key[0]} / {key[1]}: duplicate value")
        seen.add(key)
    return CheckResult("K16", "Every list is used and every value has a meaning and is unique", "MEDIUM",
                       len(lists) + len(wb.allowed), bad, "Remove unused lists; give every value a distinct meaning.")


def k17(wb: ArtifactWorkbook) -> CheckResult:
    lists = wb.list_values()
    bad = []
    total = 0
    for table, rows in wb.data.items():
        for reg in wb.register:
            if reg.get("Table") != table:
                continue
            ref = _list_ref(reg.get("Allowed values", ""))
            col = reg.get("Column")
            if not ref or ref not in lists:
                continue
            for i, row in enumerate(rows, start=1):
                value = row.get(col, "")
                if value == "":
                    continue
                total += 1
                if value not in lists[ref]:
                    bad.append(f"{DATA_SHEETS[table][0]} row {i}: {col} = '{value}' is not in list {ref}")
    return CheckResult("K17", "Recorded data stays inside the allowed lists", "HIGH", total, bad,
                       "Correct the value or add it to the list on purpose.",
                       note="Checks the NFR register, Requirements and Scope lines sheets.")


def k18(wb: ArtifactWorkbook) -> CheckResult:
    pattern = re.compile(r"^[a-z][a-z0-9_]*$")
    names = sorted({r.get("Table", "") for r in wb.register} | {f"{r.get('Table')}.{r.get('Column')}" for r in wb.register})
    bad = [n for n in names if not all(pattern.match(part) for part in n.split(".") if part)]
    return CheckResult("K18", "Table and column names are snake_case", "LOW", len(names), bad,
                       "Rename the listed items in snake_case.")


def k19(wb: ArtifactWorkbook) -> CheckResult:
    tables = sorted({r.get("Table", "") for r in wb.register if r.get("Table")})
    bad = []
    for t in tables:
        if t in RESERVED_WORDS:
            bad.append(f"table {t}: reserved word")
        elif t.endswith("s") and not t.endswith("ss") and t not in SINGULAR_EXCEPTIONS:
            bad.append(f"table {t}: looks plural")
    for r in wb.register:
        if r.get("Column", "") in RESERVED_WORDS:
            bad.append(f"column {r.get('Table')}.{r.get('Column')}: reserved word")
    return CheckResult("K19", "Table names are singular and no name is a reserved word", "LOW",
                       len(tables) + len(wb.register), bad, "Rename the listed items.")


def k20(wb: ArtifactWorkbook) -> CheckResult:
    bad = [f"{r.get('Table')}.{r.get('Column')}: {r.get('Allowed values')}" for r in wb.register
           if "comma separated" in r.get("Allowed values", "").lower()]
    return CheckResult("K20", "No column stores several values as comma separated text", "MEDIUM", len(wb.register), bad,
                       "Move each listed multi-valued column into its own child table.")


CHECKS: dict[str, Callable[[ArtifactWorkbook], CheckResult]] = {
    f"K{n:02d}": globals()[f"k{n:02d}"] for n in range(1, 21)
}


# --------------------------------------------------------------------------- scoring
def band(ratio: float) -> int:
    if ratio >= 1.0:
        return 5
    if ratio >= 0.98:
        return 4
    if ratio >= 0.90:
        return 3
    if ratio >= 0.75:
        return 2
    if ratio > 0:
        return 1
    return 0


def score_code_dimension(results: list[CheckResult]) -> int:
    ratio = sum(r.pass_rate for r in results) / len(results)
    score = band(ratio)
    if any(r.severity == "HIGH" and r.violations for r in results):
        score = min(score, 3)
    elif any(r.severity == "MEDIUM" and r.violations for r in results):
        score = min(score, 4)
    return score


def run_code_dimension(wb: ArtifactWorkbook, dim: Dimension) -> tuple[int, list[CheckResult]]:
    results = [CHECKS[cid](wb) for cid in dim.checks]
    return score_code_dimension(results), results
