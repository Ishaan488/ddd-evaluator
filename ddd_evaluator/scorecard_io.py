"""Scorecards live in Excel so mentors can write and review them without code.

Layout of a scorecard workbook:
  Scorecard    one row per setting: ID, version, what it judges, frozen, thresholds
  Dimensions   one row per dimension, with its six anchors (scores 5 down to 0)
  Instructions how to fill it in (ignored by the reader)
  Code checks  the checks a code dimension may use (ignored by the reader)

The reader returns every problem it finds in one go, with the sheet and row, so the author can fix them together.
"""
from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path
from typing import Optional

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import ValidationError

from .models import Dimension, Scorecard, Thresholds

META_SHEET = "Scorecard"
DIM_SHEET = "Dimensions"

DIM_COLUMNS = ["ID", "Name", "Group", "Weight", "Critical", "Scored by", "Measure", "Guidance", "Look for",
               "Checks", "Comes from", "Anchor 5", "Anchor 4", "Anchor 3", "Anchor 2", "Anchor 1", "Anchor 0"]
YES = {"yes", "y", "true", "1"}
NO = {"no", "n", "false", "0", ""}


class ScorecardFileError(Exception):
    """The scorecard workbook is not valid. .problems lists every issue found."""

    def __init__(self, problems: list[str]):
        self.problems = problems
        super().__init__("The scorecard has problems:\n" + "\n".join(f"- {p}" for p in problems))


def content_hash(card: Scorecard) -> str:
    """Hash of what the scorecard says, not of the file. Re-saving the workbook in Excel does not change it."""
    canonical = json.dumps(card.model_dump(), sort_keys=True, ensure_ascii=False, default=str)
    return hashlib.sha256(canonical.encode("utf-8")).hexdigest()


def _text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _lines(v) -> list[str]:
    return [p.strip() for p in _text(v).split("\n") if p.strip()]


def _norm(label: str) -> str:
    return re.sub(r"\s+", " ", label.strip().lower())


# =========================================================================== reader
def read_scorecard(path: Path) -> Scorecard:
    from .checks import CHECKS

    problems: list[str] = []
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as exc:
        raise ScorecardFileError([f"{path.name} could not be opened as an Excel workbook ({exc})."]) from exc
    missing = [s for s in (META_SHEET, DIM_SHEET) if s not in wb.sheetnames]
    if missing:
        raise ScorecardFileError([f"Sheet(s) missing: {', '.join(missing)}. Sheets found: {', '.join(wb.sheetnames)}. "
                                  "Start from the template."])

    # ---- settings
    settings: dict[str, str] = {}
    for row in wb[META_SHEET].iter_rows(values_only=True):
        if row and row[0] is not None and _text(row[0]):
            settings[_norm(_text(row[0]))] = _text(row[1] if len(row) > 1 else None)
    for key, label in (("scorecard id", "Scorecard ID"), ("version", "Version"), ("frozen", "Frozen")):
        if not settings.get(key):
            problems.append(f"Sheet '{META_SHEET}': '{label}' is empty.")
    frozen_text = settings.get("frozen", "").lower()
    if frozen_text and frozen_text not in YES | NO:
        problems.append(f"Sheet '{META_SHEET}': 'Frozen' must be yes or no, not '{settings['frozen']}'.")

    def number(key: str, label: str, default: float) -> float:
        raw = settings.get(key, "")
        if not raw:
            return default
        try:
            return float(raw)
        except ValueError:
            problems.append(f"Sheet '{META_SHEET}': '{label}' must be a number, not '{raw}'.")
            return default

    green = number("green threshold", "Green threshold", 90)
    amber = number("amber threshold", "Amber threshold", 75)
    crit_min = number("critical minimum", "Critical minimum", 4)
    if not 0 <= amber < green <= 100:
        problems.append(f"Sheet '{META_SHEET}': thresholds must satisfy 0 <= Amber < Green <= 100 (got Amber {amber:g}, Green {green:g}).")
    if not 0 <= crit_min <= 5:
        problems.append(f"Sheet '{META_SHEET}': 'Critical minimum' must be between 0 and 5.")

    # ---- dimensions table
    ws = wb[DIM_SHEET]
    rows = list(ws.iter_rows(values_only=True))
    header_idx = next((i for i, r in enumerate(rows) if r and any(_norm(_text(c)) == "id" for c in r)), None)
    if header_idx is None:
        raise ScorecardFileError(problems + [f"Sheet '{DIM_SHEET}': no header row with an 'ID' column."])
    header = [_norm(_text(c)) for c in rows[header_idx]]
    col = {name: header.index(_norm(name)) for name in DIM_COLUMNS if _norm(name) in header}
    required = ["ID", "Name", "Weight", "Scored by", "Measure"] + [f"Anchor {n}" for n in range(5, -1, -1)]
    absent = [c for c in required if c not in col]
    if absent:
        raise ScorecardFileError(problems + [f"Sheet '{DIM_SHEET}': column(s) missing: {', '.join(absent)}."])

    def cell(r, name: str):
        i = col.get(name)
        return r[i] if i is not None and i < len(r) else None

    dims: list[Dimension] = []
    for offset, r in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        if not any(_text(c) for c in r):
            continue
        dim_id = _text(cell(r, "ID"))
        where = f"Sheet '{DIM_SHEET}' row {offset}" + (f" ({dim_id})" if dim_id else "")
        row_problems: list[str] = []
        if not dim_id:
            row_problems.append("ID is empty")
        if not _text(cell(r, "Name")):
            row_problems.append("Name is empty")
        try:
            weight = float(_text(cell(r, "Weight")))
            if weight <= 0:
                row_problems.append("Weight must be above 0")
        except ValueError:
            weight = 0.0
            row_problems.append(f"Weight must be a number, not '{_text(cell(r, 'Weight'))}'")
        kind = _text(cell(r, "Scored by")).lower()
        if kind not in {"code", "llm", "human"}:
            row_problems.append(f"'Scored by' must be code, llm or human, not '{_text(cell(r, 'Scored by'))}'")
        crit_text = _text(cell(r, "Critical")).lower()
        if crit_text not in YES | NO:
            row_problems.append(f"'Critical' must be yes or no, not '{_text(cell(r, 'Critical'))}'")
        if not _text(cell(r, "Measure")):
            row_problems.append("Measure is empty (say what the dimension measures)")
        anchors = {}
        for n in range(5, -1, -1):
            text = _text(cell(r, f"Anchor {n}"))
            if not text:
                row_problems.append(f"Anchor {n} is empty")
            anchors[n] = text
        checks = [c.upper() for c in re.split(r"[,\s;]+", _text(cell(r, "Checks"))) if c]
        if kind == "code":
            if not checks:
                row_problems.append("a code dimension needs at least one check in 'Checks' (see the 'Code checks' sheet)")
            unknown = [c for c in checks if c not in CHECKS]
            if unknown:
                row_problems.append(f"unknown check(s) {', '.join(unknown)}; valid ids are {', '.join(sorted(CHECKS))}")
        elif checks:
            row_problems.append("'Checks' is only for code dimensions; leave it empty for llm and human")
        if kind == "llm" and not _text(cell(r, "Guidance")):
            row_problems.append("an llm dimension needs Guidance, so every rater reads the same instruction")
        if row_problems:
            problems.extend(f"{where}: {p}." for p in row_problems)
            continue
        dims.append(Dimension(
            id=dim_id, name=_text(cell(r, "Name")), group=_text(cell(r, "Group")), weight=weight,
            critical=crit_text in YES, evaluator=kind, measure=_text(cell(r, "Measure")),
            guidance=_text(cell(r, "Guidance")), look_for=_lines(cell(r, "Look for")), checks=checks,
            traces_to=_lines(cell(r, "Comes from")), anchors=anchors,
        ))
    if not dims and not problems:
        problems.append(f"Sheet '{DIM_SHEET}': no dimensions found below the header row.")
    ids = [d.id for d in dims]
    for dup in sorted({i for i in ids if ids.count(i) > 1}):
        problems.append(f"Sheet '{DIM_SHEET}': ID '{dup}' is used more than once.")
    if dims and not problems:
        total = sum(d.weight for d in dims)
        if abs(total - 100) > 0.01:
            problems.append(f"Sheet '{DIM_SHEET}': weights add up to {total:g}, not 100.")
    if problems:
        raise ScorecardFileError(problems)

    try:
        return Scorecard(
            id=settings["scorecard id"], version=settings["version"], judges=settings.get("judges", ""),
            frozen=frozen_text in YES, thresholds=Thresholds(green=green, amber=amber, critical_min=int(crit_min)),
            dimensions=dims,
        )
    except ValidationError as exc:  # should be unreachable after the checks above
        raise ScorecardFileError([str(e["msg"]) for e in exc.errors()]) from exc


# =========================================================================== writer
HEAD_FILL = PatternFill("solid", fgColor="1D4ED8")
HEAD_FONT = Font(bold=True, color="FFFFFF")
WRAP = Alignment(wrap_text=True, vertical="top")
WIDTHS = {"ID": 8, "Name": 30, "Group": 18, "Weight": 9, "Critical": 10, "Scored by": 11, "Measure": 42,
          "Guidance": 60, "Look for": 42, "Checks": 18, "Comes from": 36}

INSTRUCTIONS = [
    ("How to fill in this scorecard", None),
    ("1", "Sheet 'Scorecard': give the scorecard an ID and a version, say what it judges, and set the thresholds. "
          "Set Frozen to yes only when the scorecard is final. The evaluator refuses a scorecard that is not frozen."),
    ("2", "Sheet 'Dimensions': one row per thing you want to score. Weights must add up to 100."),
    ("3", "Scored by: 'code' uses built-in checks (list them in the Checks column; see sheet 'Code checks'). 'llm' is scored 0 to 5 "
          "by three independent models, so write Guidance that tells them exactly what to look at. 'human' is scored by a person "
          "and supplied when the evaluation is run; the agent never scores it."),
    ("4", "Anchors: write what a 5, 4, 3, 2, 1 and 0 each look like, in words a rater can check against the artifact. "
          "Vague anchors give inconsistent scores."),
    ("5", "Critical: yes means the dimension must score at least the Critical minimum, or the verdict cannot be green."),
    ("6", "Look for and Comes from: one item per line (press Alt+Enter inside a cell for a new line). Optional."),
    ("", ""),
    ("How the verdict is decided", None),
    ("Green", "Weighted score at or above the Green threshold, every critical dimension at or above the Critical minimum, "
              "no high-risk finding open, and every human dimension scored."),
    ("Amber", "Weighted score between Amber and Green, or a critical dimension below the minimum, or a high-risk finding open."),
    ("Red", "Weighted score below the Amber threshold."),
    ("", ""),
    ("Freeze it before you use it", "Decide the scorecard before you see the artifact it judges, so the target does not move. "
                                    "Record why you chose each dimension in your decision log."),
]


def _style_header(ws, widths: dict[str, float]) -> None:
    for c in ws[1]:
        c.fill, c.font, c.alignment = HEAD_FILL, HEAD_FONT, Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[c.column_letter].width = widths.get(str(c.value), 38 if str(c.value).startswith("Anchor") else 16)
    ws.freeze_panes = "C2"


def write_scorecard_xlsx(card: Scorecard, path: Path) -> None:
    from .checks import ArtifactWorkbook, CHECKS

    wb = openpyxl.Workbook()
    meta = wb.active
    meta.title = META_SHEET
    meta.append(["Setting", "Value", "Note"])
    for c in meta[1]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
    rows = [
        ("Scorecard ID", card.id, "A short unique name, for example SC-STEP03-D2."),
        ("Version", card.version, "Text, for example 1.0."),
        ("Judges", card.judges, "What artifact this scorecard judges."),
        ("Frozen", "yes" if card.frozen else "no", "yes or no. Only a frozen scorecard can drive a verdict."),
        ("Green threshold", card.thresholds.green, "Weighted score needed for green."),
        ("Amber threshold", card.thresholds.amber, "Below this the verdict is red."),
        ("Critical minimum", card.thresholds.critical_min, "Lowest score a critical dimension may have (0 to 5)."),
    ]
    for r in rows:
        meta.append(list(r))
    for row in meta.iter_rows(min_row=2, max_col=2):
        row[1].number_format = "@"
        row[1].alignment = WRAP
    meta.append([])
    meta.append(["Weights add up to (must be 100)", f"=SUM({DIM_SHEET}!D2:D500)", "Live total of the Weight column."])
    for col, w in zip("ABC", (32, 60, 60)):
        meta.column_dimensions[col].width = w
    dv = DataValidation(type="list", formula1='"yes,no"', allow_blank=False)
    meta.add_data_validation(dv)
    dv.add("B5")

    ws = wb.create_sheet(DIM_SHEET)
    ws.append(DIM_COLUMNS)
    for d in card.dimensions:
        ws.append([d.id, d.name, d.group, d.weight, "yes" if d.critical else "no", d.evaluator, d.measure, d.guidance,
                   "\n".join(d.look_for), ", ".join(d.checks), "\n".join(d.traces_to)]
                  + [d.anchors[n] for n in range(5, -1, -1)])
    _style_header(ws, WIDTHS)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
    last = max(ws.max_row, 2)
    for formula, rng in (('"yes,no"', f"E2:E{last + 40}"), ('"code,llm,human"', f"F2:F{last + 40}")):
        v = DataValidation(type="list", formula1=formula, allow_blank=True)
        ws.add_data_validation(v)
        v.add(rng)

    ins = wb.create_sheet("Instructions")
    for a, b in INSTRUCTIONS:
        ins.append([a, b])
    for row in ins.iter_rows():
        row[0].font = Font(bold=True)
        row[0].alignment = Alignment(vertical="top", wrap_text=True)
        row[1].alignment = WRAP
        if row[1].value is None and row[0].value:
            row[0].font = Font(bold=True, size=13)
    ins.column_dimensions["A"].width = 30
    ins.column_dimensions["B"].width = 110

    chk = wb.create_sheet("Code checks")
    chk.append(["Check", "What it tests", "Severity"])
    empty = ArtifactWorkbook([], [], [], {})
    for cid in sorted(CHECKS):
        r = CHECKS[cid](empty)
        chk.append([cid, r.name, r.severity])
    for c in chk[1]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
    chk.column_dimensions["A"].width, chk.column_dimensions["B"].width, chk.column_dimensions["C"].width = 10, 80, 10
    chk.append([])
    chk.append(["", "These checks read a workbook with sheets 'Tree', 'Schema register' and 'Allowed values'. "
                    "On any other artifact a code dimension reports 'not measurable'.", ""])
    wb.save(path)


def template_card() -> Scorecard:
    anchors = {5: "Describe what a perfect result looks like.", 4: "Describe a strong result with one small gap.",
               3: "Describe an acceptable result with clear gaps.", 2: "Describe a weak result.",
               1: "Describe a very weak result.", 0: "Describe no result or a missing deliverable."}
    return Scorecard(
        id="SC-EXAMPLE", version="0.1", judges="Describe the artifact this scorecard judges", frozen=False,
        dimensions=[
            Dimension(id="E01", name="Example code dimension (delete me)", group="Structure", weight=20, critical=True,
                      evaluator="code", measure="Share of the listed checks that pass.", checks=["K01", "K02", "K03"], anchors=anchors),
            Dimension(id="E02", name="Example LLM dimension (delete me)", group="Quality", weight=60, critical=False,
                      evaluator="llm", measure="What the raters should measure.",
                      guidance="Tell the raters exactly what to read and what a good answer contains.",
                      look_for=["One thing to look for", "Another thing to look for"], anchors=anchors),
            Dimension(id="E03", name="Example human dimension (delete me)", group="Human gates", weight=20, critical=False,
                      evaluator="human", measure="A person scores this, for example in an interview.", anchors=anchors),
        ],
    )


def write_template(path: Path) -> None:
    write_scorecard_xlsx(template_card(), path)
