"""Load and freeze the inputs: artifact, scorecard, source documents, config."""
from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

import yaml

from .models import Scorecard  # noqa: F401  (returned by load_scorecard)

TEXT_SUFFIXES = {".md", ".txt", ".sql", ".yaml", ".yml", ".json", ".csv", ".dbml"}
XLSX_SUFFIXES = {".xlsx", ".xlsm"}
DEFAULT_ARTIFACT_SHEETS = ["Tree", "Schema register", "Allowed values"]


class InputError(Exception):
    """A problem with an input the user can fix."""


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: Path) -> str:
    return sha256_bytes(path.read_bytes())


@dataclass
class LoadedDoc:
    spec: str
    path: Path
    label: str
    kind: str
    text: str
    sha256: str
    sheets: list[str] = field(default_factory=list)


def split_spec(spec: str) -> tuple[Path, Optional[list[str]]]:
    """'file.xlsx::Sheet A,Sheet B' -> (path, ['Sheet A', 'Sheet B']). Windows drive colons are safe."""
    if "::" in spec:
        raw_path, raw_sheets = spec.rsplit("::", 1)
        sheets = [s.strip() for s in raw_sheets.split(",") if s.strip()]
        return Path(raw_path), sheets or None
    return Path(spec), None


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).replace("\r", " ").replace("\n", " ").strip()


def xlsx_to_text(path: Path, sheets: Optional[list[str]]) -> tuple[str, list[str]]:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    names = wb.sheetnames
    if sheets is None:
        chosen = names
    else:
        missing = [s for s in sheets if s not in names]
        if missing:
            raise InputError(f"{path.name}: no sheet named {missing}. Sheets present: {names}")
        chosen = sheets
    parts: list[str] = []
    for name in chosen:
        ws = wb[name]
        parts.append(f"### SHEET: {name}")
        for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            cells = [_cell(c) for c in row]
            while cells and cells[-1] == "":
                cells.pop()
            if not cells:
                continue
            parts.append(f"R{idx}: " + " | ".join(cells))
        parts.append("")
    return "\n".join(parts), chosen


def pdf_to_text(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - depends on environment
        raise InputError(
            f"{path.name} is a PDF and reading PDFs needs pypdf. Run: pip install pypdf "
            "(or give the docket as .md or .txt)."
        ) from exc
    reader = PdfReader(str(path))
    return "\n\n".join(f"[page {i}]\n{(p.extract_text() or '').strip()}" for i, p in enumerate(reader.pages, 1))


def load_doc(spec: str, *, default_sheets: Optional[list[str]] = None) -> LoadedDoc:
    path, sheets = split_spec(spec)
    if not path.is_file():
        raise InputError(f"file not found: {path}")
    suffix = path.suffix.lower()
    raw_hash = sha256_file(path)
    used_sheets: list[str] = []
    if suffix in XLSX_SUFFIXES:
        wanted = sheets
        if wanted is None and default_sheets:
            import openpyxl

            present = openpyxl.load_workbook(path, read_only=True).sheetnames
            if all(s in present for s in default_sheets):
                wanted = default_sheets
        text, used_sheets = xlsx_to_text(path, wanted)
        kind = "xlsx"
    elif suffix == ".pdf":
        text, kind = pdf_to_text(path), "pdf"
    elif suffix in TEXT_SUFFIXES:
        text, kind = path.read_text(encoding="utf-8", errors="replace"), "text"
    else:
        raise InputError(f"unsupported file type {suffix!r} for {path.name}. Use xlsx, pdf, md, txt, sql, yaml, json or csv.")
    label = path.name + (f" [{', '.join(used_sheets)}]" if used_sheets else "")
    return LoadedDoc(spec=spec, path=path, label=label, kind=kind, text=text, sha256=raw_hash, sheets=used_sheets)


def detect_version(text: str, path: Optional[Path] = None) -> str:
    """Find 'Schema v1.4' in the artifact text, or in the Summary sheet of the workbook it came from."""
    pattern = re.compile(r"[Ss]chema\s+v(\d+(?:\.\d+)+)")
    m = pattern.search(text)
    if m or path is None or path.suffix.lower() not in XLSX_SUFFIXES:
        return m.group(1) if m else ""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if "Summary" in wb.sheetnames:
        for row in wb["Summary"].iter_rows(values_only=True):
            for cell in row:
                m = pattern.search(str(cell)) if cell else None
                if m:
                    return m.group(1)
    return ""


# --------------------------------------------------------------------------- scorecard and config
def load_scorecard(path: Path, *, allow_unfrozen: bool = False) -> tuple[Scorecard, str]:
    """Read a scorecard workbook. Returns the scorecard and a hash of its content (not of the file)."""
    from .scorecard_io import ScorecardFileError, content_hash, read_scorecard

    if not path.is_file():
        raise InputError(f"scorecard not found: {path}")
    if path.suffix.lower() not in XLSX_SUFFIXES:
        raise InputError(f"{path.name}: scorecards are Excel files (.xlsx). "
                         "Create one with: python -m ddd_evaluator template --out scorecard.xlsx")
    try:
        card = read_scorecard(path)
    except ScorecardFileError as exc:
        raise InputError(str(exc)) from exc
    if not card.frozen and not allow_unfrozen:
        raise InputError(
            f"scorecard {card.id} v{card.version} is not frozen. A verdict must not be driven by a scorecard that can "
            "still change. Set 'Frozen' to yes in the 'Scorecard' sheet, or pass --allow-unfrozen for a trial run."
        )
    return card, content_hash(card)


def load_env_file(path: Path = Path(".env")) -> None:
    """Tiny .env reader so no extra dependency is needed. Existing environment wins."""
    import os

    if not path.is_file():
        return
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, _, value = line.partition("=")
        os.environ.setdefault(key.strip(), value.strip().strip('"').strip("'"))


def load_config(path: Path) -> dict:
    if not path.is_file():
        raise InputError(f"config not found: {path}")
    return yaml.safe_load(path.read_text(encoding="utf-8"))


def load_human_evidence(path: Optional[Path]) -> dict[str, dict]:
    if path is None:
        return {}
    if not path.is_file():
        raise InputError(f"human evidence file not found: {path}")
    data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    out: dict[str, dict] = {}
    for dim_id, item in data.items():
        if not isinstance(item, dict) or "score" not in item:
            raise InputError(f"human evidence for {dim_id} needs a 'score' (0 to 5), plus 'by' and 'note'")
        score = int(item["score"])
        if not 0 <= score <= 5:
            raise InputError(f"human evidence score for {dim_id} must be 0 to 5")
        if not str(item.get("by", "")).strip():
            raise InputError(f"human evidence for {dim_id} must say who supplied it ('by')")
        out[str(dim_id)] = {"score": score, "by": str(item["by"]), "note": str(item.get("note", ""))}
    return out
